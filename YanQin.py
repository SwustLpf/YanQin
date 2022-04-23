import docx
import xlsxwriter
import copy
import re
import sys


def getHanziNum(strCSrc):
    hanzi = re.findall("[\u4e00-\u9fa5]",strCSrc)
    hanzi = "".join(hanzi)
    num = re.findall(r"\d+",strCSrc)
    return [hanzi,num]

class CCategory():
    def __init__(self):
        self.m_Name = ""
        self.m_Supply = []

    def parseStr(self,strCSrc):
        str1 = strCSrc.split("：")
        if(1 >= len(str1)):
            str1 = strCSrc.split(":")
            
        if(2 <= len(str1)):
            self.m_Name = str1[0]
            str2 = str1[1]
            str3 = str2.split()
            self.m_Supply = str3

    def getName(self):
        return self.m_Name

    def getSupply(self):
        return self.m_Supply

class CTaocan():
    def __init__(self):
        self.m_Name = ""
        self.m_Huncai = []
        self.m_Sucai = []
        self.m_Zhushi = []
        self.m_Shuiguo = []
        self.m_Naizhipin = []
        self.m_Xiafancai = []
        self.m_Baipan = []
    
    def putName(self,strCName):
        self.m_Name = strCName
    def getName(self):
        return self.m_Name

    def addHuncai(self,huicai):
        self.m_Huncai.append(huicai)
    def getHuncai(self):
        return self.m_Huncai
    
    def addSucai(self,sucai):
        self.m_Sucai.append(sucai)
    def getSucai(self):
        return self.m_Sucai

    def addZhushi(self,zhushi):
        self.m_Zhushi.append(zhushi)
    def getZhushi(self):
        return self.m_Zhushi

    def addShuiguo(self,sucai):
        self.m_Shuiguo.append(sucai)
    def getShuiguo(self):
        return self.m_Shuiguo

    def addNaizhipin(self,sucai):
        self.m_Naizhipin.append(sucai)
    def getNaizhipin(self):
        return self.m_Naizhipin

    def addXiafancai(self,sucai):
        self.m_Xiafancai.append(sucai)
    def getXiafancai(self):
        return self.m_Xiafancai

    def addBaipan(self,sucai):
        self.m_Baipan.append(sucai)
    def getBaipan(self):
        return self.m_Baipan

    def printInfo(self):
        print(self.m_Name)
        print("下面是荤菜")
        for huncai in self.m_Huncai:
            print(huncai.getName(),huncai.getSupply())
        print("下面是素菜")
        for sucai in self.m_Sucai:
            print(sucai.getName(),sucai.getSupply())
    def clearAll(self):
        self.m_Name = ""
        self.m_Huncai.clear()
        self.m_Sucai.clear()
        # self.m_Huncai = []
        # self.m_Sucai = []

s_Taocan = CTaocan()
s_Ans = []

s_taocan = ""
s_huncai = 0 #荤菜
s_sucai = 0 #素菜
s_zhushi = "" #主食
s_shuiguo = "" #水果
s_naizhiping = ""
s_xiafancai = ""
s_baipan = ""

s_strHuncai = []
s_strSucai = []

def maohao(strCSrc):
    strResult = strCSrc.split("：")
    if(1 >= len(strResult)):
        strResult = strCSrc.split(":")
    return strResult[1]

def f(strCSrc):
    # print("strCSrc = ",strCSrc)
    str1 = strCSrc.split("：")
    if(1 >= len(str1)):
        str1 = strCSrc.split(":")
        
    if(1 >= len(str1)):
        return []
    # print("len = ",len(str1))
    # print("str1 = ",str1)
    strTitle = str1[0]
    
    str2 = str1[1]
    # print("str2 = ",str2)
    str3 = str2.split("g")
    str3 = str3[0:len(str3)-1]
    # print("str3 = ",str3)
    return [strTitle,str3]


if __name__ == '__main__':
    

    if(len(sys.argv)<3):
        print("请输入源文件名 与 目标文件名")
        print("示例: python YanQin.py 机组正餐.docx 机组餐测算.xlsx")
        exit(1)
        
    #获取文档对象
    file=docx.Document(sys.argv[1])


    for para in file.paragraphs:
        strLineData = para.text
        if("XX航空机组正餐" == strLineData):
            continue

        # print("strLineData = ",strLineData)
        if(len(strLineData) < 1):
            continue
        if(strLineData[0].encode().isalpha()):
            # print("strLineData = ",strLineData)
            if(len(s_Taocan.getName())>0):
                s_Ans.append(s_Taocan)
            # s_Taocan.printInfo()

            TmpTaocan = CTaocan()
            s_Taocan = copy.deepcopy(TmpTaocan)

            s_Taocan.putName(strLineData)
            s_huncai = 0
            s_sucai = 0
        elif ("主食" == strLineData[0:2]):
            objcai = CCategory()
            objcai.parseStr(strLineData)
            s_Taocan.addZhushi(objcai)
        elif ("水果" == strLineData[0:2]):
            objcai = CCategory()
            objcai.parseStr(strLineData)
            s_Taocan.addShuiguo(objcai)
        elif ("奶制品" == strLineData[0:3]):
            objcai = CCategory()
            objcai.parseStr(strLineData)
            s_Taocan.addNaizhipin(objcai)
        elif ("下饭菜" == strLineData[0:3]):
            objcai = CCategory()
            objcai.parseStr(strLineData)
            s_Taocan.addXiafancai(objcai)
        elif ("摆盘" == strLineData[0:2]):
            objcai = CCategory()
            objcai.parseStr(strLineData)
            s_Taocan.addBaipan(objcai)
        elif ("荤菜" == strLineData[0:2]):
            s_huncai = 1
            s_sucai = 0
        elif ("素菜" == strLineData[0:2]):
            s_huncai = 0
            s_sucai = 1
        else:
            objcai = CCategory()
            objcai.parseStr(strLineData)
            if(1 == s_huncai):
                s_Taocan.addHuncai(objcai)
            else:
                s_Taocan.addSucai(objcai)
    s_Ans.append(s_Taocan) # 最后一个数据


    

    workbook = xlsxwriter.Workbook(sys.argv[2]) 
    worksheet = workbook.add_worksheet("sheet1") 
    worksheet.merge_range(0,0,0,6, 'XX航空机组正餐成本明细')
    worksheet.write("A2","类别")
    worksheet.write("B2","序号")
    worksheet.write("C2","餐食名称")
    worksheet.write("D2","材料")
    worksheet.write("E2","重量(g)")
    worksheet.write("F2","单价(元/g）")
    worksheet.write("G2","税率")

    # s_Taocan.printInfo()


    # for j in range(0,len(s_Ans)):
    #             print("j = ",j)
    #             s_Ans[j].printInfo()

    # print("len(s_Ans) = ",len(s_Ans))
    #写入荤菜
    s_row = 2
    startHun = s_row
    for t in s_Ans:
        # t.printInfo()
        startRowTaocan = s_row
        for cai in t.getHuncai():
            startRow = s_row
            for yuanliao in cai.getSupply():
                Tmp = getHanziNum(yuanliao)
                worksheet.write(s_row,3,Tmp[0])
                worksheet.write(s_row,4,"".join(Tmp[1])+yuanliao[-1])
                s_row = s_row + 1
            worksheet.merge_range(startRow,2,s_row-1,2, cai.getName())
        worksheet.merge_range(startRowTaocan,1,s_row-1,1, t.getName())
    worksheet.merge_range(startHun,0,s_row-1,0, "荤菜")



    #写入素菜
    startSu = s_row
    for t in s_Ans:
        # t.printInfo()
        startRowTaocan = s_row
        for cai in t.getSucai():
            startRow = s_row
            for yuanliao in cai.getSupply():
                Tmp = getHanziNum(yuanliao)
                worksheet.write(s_row,3,Tmp[0])
                worksheet.write(s_row,4,"".join(Tmp[1])+yuanliao[-1])
                s_row = s_row + 1
            worksheet.merge_range(startRow,2,s_row-1,2, cai.getName())
        worksheet.merge_range(startRowTaocan,1,s_row-1,1, t.getName())
    worksheet.merge_range(startSu,0,s_row-1,0, "素菜")


    #写入主食
    startZhushi = s_row
    for t in s_Ans:
        # t.printInfo()
        startRowTaocan = s_row
        for cai in t.getZhushi():
            startRow = s_row
            for yuanliao in cai.getSupply():
                Tmp = getHanziNum(yuanliao)
                # print("Tmp[0] = ",Tmp[0])
                # print("Tmp[1] = ",Tmp[1])
                worksheet.write(s_row,3,Tmp[0])
                worksheet.write(s_row,4,"".join(Tmp[1])+yuanliao[-1])
                s_row = s_row + 1
            if(s_row-1>startRow):
                worksheet.merge_range(startRow,2,s_row-1,2, cai.getName())
            else:
                worksheet.write(startRow,2,cai.getName())
        if(s_row-1>startRowTaocan): 
            worksheet.merge_range(startRowTaocan,1,s_row-1,1, t.getName())
        else:
            worksheet.write(startRowTaocan,1,t.getName())
    # print("startZhushi = ",startZhushi)
    # print("s_row-1 = ",s_row-1)
    if(s_row-1>startZhushi): 
        worksheet.merge_range(startZhushi,0,s_row-1,0, "主食")
    else:
        worksheet.write(startZhushi,0,"主食")


    #写入水果
    startShuiguo = s_row
    for t in s_Ans:
        # t.printInfo()
        startRowTaocan = s_row
        for cai in t.getShuiguo():
            startRow = s_row
            for yuanliao in cai.getSupply():
                Tmp = getHanziNum(yuanliao)
                # print("Tmp[0] = ",Tmp[0])
                # print("Tmp[1] = ",Tmp[1])
                worksheet.write(s_row,3,Tmp[0])
                worksheet.write(s_row,4,"".join(Tmp[1])+yuanliao[-1])
                s_row = s_row + 1
            if(s_row-1>startRow):
                worksheet.merge_range(startRow,2,s_row-1,2, cai.getName())
            else:
                worksheet.write(startRow,2,cai.getName())
        if(s_row-1>startRowTaocan): 
            worksheet.merge_range(startRowTaocan,1,s_row-1,1, t.getName())
        else:
            worksheet.write(startRowTaocan,1,t.getName())
    # print("startZhushi = ",startZhushi)
    # print("s_row-1 = ",s_row-1)
    if(s_row-1>startShuiguo): 
        worksheet.merge_range(startShuiguo,0,s_row-1,0, "水果")
    else:
        worksheet.write(startShuiguo,0,"水果")

    #写入奶制品
    startNaizhipin = s_row
    for t in s_Ans:
        # t.printInfo()
        startRowTaocan = s_row
        for cai in t.getNaizhipin():
            startRow = s_row
            for yuanliao in cai.getSupply():
                Tmp = getHanziNum(yuanliao)
                # print("Tmp[0] = ",Tmp[0])
                # print("Tmp[1] = ",Tmp[1])
                worksheet.write(s_row,3,Tmp[0])
                worksheet.write(s_row,4,"".join(Tmp[1])+yuanliao[-1])
                s_row = s_row + 1
            if(s_row-1>startRow):
                worksheet.merge_range(startRow,2,s_row-1,2, cai.getName())
            else:
                worksheet.write(startRow,2,cai.getName())
        if(s_row-1>startRowTaocan): 
            worksheet.merge_range(startRowTaocan,1,s_row-1,1, t.getName())
        else:
            worksheet.write(startRowTaocan,1,t.getName())
    # print("startZhushi = ",startZhushi)
    # print("s_row-1 = ",s_row-1)
    if(s_row-1>startNaizhipin): 
        worksheet.merge_range(startNaizhipin,0,s_row-1,0, "奶制品")
    else:
        worksheet.write(startNaizhipin,0,"奶制品")

    #写入下饭菜
    startXiafancai = s_row
    for t in s_Ans:
        # t.printInfo()
        startRowTaocan = s_row
        for cai in t.getXiafancai():
            startRow = s_row
            for yuanliao in cai.getSupply():
                Tmp = getHanziNum(yuanliao)
                # print("Tmp[0] = ",Tmp[0])
                # print("Tmp[1] = ",Tmp[1])
                worksheet.write(s_row,3,Tmp[0])
                worksheet.write(s_row,4,"".join(Tmp[1])+yuanliao[-1])
                s_row = s_row + 1
            if(s_row-1>startRow):
                worksheet.merge_range(startRow,2,s_row-1,2, cai.getName())
            else:
                worksheet.write(startRow,2,cai.getName())
        if(s_row-1>startRowTaocan): 
            worksheet.merge_range(startRowTaocan,1,s_row-1,1, t.getName())
        else:
            worksheet.write(startRowTaocan,1,t.getName())
    # print("startZhushi = ",startZhushi)
    # print("s_row-1 = ",s_row-1)
    if(s_row-1>startXiafancai): 
        worksheet.merge_range(startXiafancai,0,s_row-1,0, "下饭菜")
    else:
        worksheet.write(startXiafancai,0,"下饭菜")

    #写入摆盘
    startBaipan = s_row
    for t in s_Ans:
        # t.printInfo()
        startRowTaocan = s_row
        for cai in t.getBaipan():
            startRow = s_row
            for yuanliao in cai.getSupply():
                Tmp = getHanziNum(yuanliao)
                # print("Tmp[0] = ",Tmp[0])
                # print("Tmp[1] = ",Tmp[1])
                worksheet.write(s_row,3,Tmp[0])
                worksheet.write(s_row,4,"".join(Tmp[1])+yuanliao[-1])
                s_row = s_row + 1
            if(s_row-1>startRow):
                worksheet.merge_range(startRow,2,s_row-1,2, cai.getName())
            else:
                worksheet.write(startRow,2,cai.getName())
        if(s_row-1>startRowTaocan): 
            worksheet.merge_range(startRowTaocan,1,s_row-1,1, t.getName())
        else:
            worksheet.write(startRowTaocan,1,t.getName())
    # print("startZhushi = ",startZhushi)
    # print("s_row-1 = ",s_row-1)
    if(s_row-1>startBaipan): 
        worksheet.merge_range(startBaipan,0,s_row-1,0, "摆盘")
    else:
        worksheet.write(startBaipan,0,"摆盘")



        # print("t.getName() = ",t.getName())
    workbook.close()



